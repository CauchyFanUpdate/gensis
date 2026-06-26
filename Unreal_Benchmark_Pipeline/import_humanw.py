#!/usr/bin/env python3
"""Import human_correct.xlsx into manual_review.jsonl.

The xlsx contains four human-annotation columns:
  - 图片有无崩坏           → humanw_image_broken     ("有"=True, "无"=False)
  - 问题是否合理           → humanw_question_ok      ("合理"=True, "不合理"=False)
  - correct_label是否正确  → humanw_correct_label_ok ("正确"=True, "不正确"=False)
  - 是否保留               → humanw_keep             ("保留"=True, "不保留"=False)

Rows are matched back to samples by the filename portion of the `image_path` URL
(e.g. .../for_humanw/ATTR-073-mc.png → sample_id = ATTR-073-mc).

For each matched sample we read its LATEST manual_review record, merge in the
four humanw_* fields, set source="humanw_import", and APPEND the merged record.
`human_verified` is NOT set to True — humanw is a first-pass annotation and the
project owner will still verify each sample manually. The original file is
backed up to manual_review.jsonl.bak.<unix-ts> before any writing.
"""
from __future__ import annotations

import json
import shutil
import sys
import time
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parent
XLSX = ROOT / "human_correct.xlsx"
MANUAL = ROOT / "manual_review.jsonl"

COL_IMAGE_PATH = "image_path"
COL_BROKEN = "图片有无崩坏"
COL_QUESTION_OK = "问题是否合理"
COL_LABEL_OK = "correct_label\n是否正确"
COL_KEEP = "是否保留（qwen和gemini答案都正确，视为不保留）"

YES_VALUES = {"有", "合理", "正确", "保留"}
NO_VALUES = {"无", "不合理", "不正确", "不保留"}


def parse_yesno(v) -> bool | None:
    if v is None:
        return None
    s = str(v).strip()
    if not s:
        return None
    if s in YES_VALUES:
        return True
    if s in NO_VALUES:
        return False
    return None


def url_to_sample_id(url: str) -> str | None:
    if not url:
        return None
    fname = url.rsplit("/", 1)[-1]
    if fname.lower().endswith(".png"):
        fname = fname[:-4]
    return fname or None


def load_latest_manual() -> dict[str, dict]:
    out: dict[str, dict] = {}
    if not MANUAL.exists():
        return out
    with MANUAL.open("r", encoding="utf-8") as f:
        for line in f:
            line = line.strip()
            if not line:
                continue
            try:
                r = json.loads(line)
            except json.JSONDecodeError:
                continue
            sid = r.get("sample_id")
            if sid:
                out[sid] = r
    return out


def read_xlsx_rows() -> list[dict]:
    wb = openpyxl.load_workbook(XLSX, data_only=True)
    ws = wb.active
    headers = [c.value for c in ws[1]]
    needed = [COL_IMAGE_PATH, COL_BROKEN, COL_QUESTION_OK, COL_LABEL_OK, COL_KEEP]
    idx = {}
    for col in needed:
        if col not in headers:
            raise SystemExit(f"missing column in xlsx: {col!r}")
        idx[col] = headers.index(col)

    rows: list[dict] = []
    for row in ws.iter_rows(min_row=2, values_only=True):
        if row[idx[COL_IMAGE_PATH]] is None:
            continue
        rows.append({
            "image_url": row[idx[COL_IMAGE_PATH]],
            "broken": row[idx[COL_BROKEN]],
            "question_ok": row[idx[COL_QUESTION_OK]],
            "label_ok": row[idx[COL_LABEL_OK]],
            "keep": row[idx[COL_KEEP]],
        })
    return rows


def main() -> None:
    if not XLSX.exists():
        raise SystemExit(f"xlsx not found: {XLSX}")
    if not MANUAL.exists():
        raise SystemExit(f"manual_review.jsonl not found: {MANUAL}")

    ts = int(time.time())
    bak = MANUAL.with_suffix(MANUAL.suffix + f".bak.{ts}")
    shutil.copy2(MANUAL, bak)
    print(f"[backup] {MANUAL.name} → {bak.name}", flush=True)

    rows = read_xlsx_rows()
    print(f"[xlsx] {len(rows)} rows", flush=True)

    latest = load_latest_manual()
    print(f"[manual] {len(latest)} unique sample_ids loaded", flush=True)

    matched: list[dict] = []
    unmatched: list[str] = []
    skipped_empty: list[str] = []
    for row in rows:
        sid = url_to_sample_id(str(row["image_url"]))
        if not sid:
            continue
        base = latest.get(sid)
        if base is None:
            unmatched.append(sid)
            continue
        broken = parse_yesno(row["broken"])
        q_ok = parse_yesno(row["question_ok"])
        l_ok = parse_yesno(row["label_ok"])
        keep = parse_yesno(row["keep"])
        if broken is None and q_ok is None and l_ok is None and keep is None:
            skipped_empty.append(sid)
            continue
        new_rec = json.loads(json.dumps(base))  # deep copy
        new_rec["humanw_image_broken"] = broken
        new_rec["humanw_question_ok"] = q_ok
        new_rec["humanw_correct_label_ok"] = l_ok
        new_rec["humanw_keep"] = keep
        new_rec["edited_at"] = ts
        new_rec["source"] = "humanw_import"
        matched.append(new_rec)

    if unmatched:
        print(f"[warn] {len(unmatched)} rows did not match any sample_id:", flush=True)
        for sid in unmatched[:10]:
            print(f"  - {sid}", flush=True)
        if len(unmatched) > 10:
            print(f"  ... +{len(unmatched)-10} more", flush=True)
    if skipped_empty:
        print(f"[info] {len(skipped_empty)} rows had no annotations, skipped", flush=True)

    if not matched:
        print("[done] nothing to append", flush=True)
        return

    with MANUAL.open("a", encoding="utf-8") as f:
        for r in matched:
            f.write(json.dumps(r, ensure_ascii=False) + "\n")
    print(f"[done] appended {len(matched)} records to {MANUAL.name}", flush=True)


if __name__ == "__main__":
    main()
